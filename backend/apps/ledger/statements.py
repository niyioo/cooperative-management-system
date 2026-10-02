"""
Member statements (Excel and PDF). A statement lists posted entries in the
period, including reversals and the entries they cancel, so the totals agree
with the member's balances. Pending and rejected entries are not part of a statement.
"""
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.configuration.models import CooperativeSettings

from .choices import EntrySide, TransactionStatus
from .selectors import member_transactions
from .serializers import TransactionSerializer

NAVY = "293C9C"


def statement_rows(member, date_from=None, date_to=None):
    entries = member_transactions(member).filter(status__in=[TransactionStatus.POSTED, TransactionStatus.REVERSED])
    if date_from:
        entries = entries.filter(value_date__gte=date_from)
    if date_to:
        entries = entries.filter(value_date__lte=date_to)
    rows = []
    for entry in entries.order_by("value_date", "created_at"):
        data = TransactionSerializer(entry).data
        account = data["account"]
        rows.append(
            {
                "date": entry.value_date,
                "reference": entry.reference,
                "description": entry.description or data["type_label"],
                "type": data["type_label"],
                "account": f"{account['label']} ({account['number']})" if account else "Paid outside the cooperative",
                "credit": entry.amount if entry.entry_side == EntrySide.CREDIT else None,
                "debit": entry.amount if entry.entry_side == EntrySide.DEBIT else None,
                "status": entry.get_status_display(),
            }
        )
    return rows


def _period_label(date_from, date_to):
    if date_from and date_to:
        return f"{date_from:%d %b %Y} to {date_to:%d %b %Y}"
    if date_from:
        return f"From {date_from:%d %b %Y}"
    if date_to:
        return f"Up to {date_to:%d %b %Y}"
    return "All transactions"


def _totals(rows):
    credits = sum((r["credit"] for r in rows if r["credit"]), Decimal("0.00"))
    debits = sum((r["debit"] for r in rows if r["debit"]), Decimal("0.00"))
    return credits, debits


def build_xlsx(member, date_from=None, date_to=None):
    coop = CooperativeSettings.load()
    rows = statement_rows(member, date_from, date_to)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Statement"
    sheet.append([coop.name])
    sheet["A1"].font = Font(bold=True, size=14, color=NAVY)
    sheet.append([f"Statement for {member.full_name} ({member.membership_number})"])
    sheet.append([_period_label(date_from, date_to)])
    sheet.append([f"Generated {timezone.localtime():%d %b %Y %H:%M}"])
    sheet.append([])
    headers = ["Date", "Reference", "Description", "Account", "Credit (₦)", "Debit (₦)", "Status"]
    sheet.append(headers)
    for cell in sheet[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for r in rows:
        sheet.append([r["date"], r["reference"], r["description"], r["account"], r["credit"], r["debit"], r["status"]])
    credits, debits = _totals(rows)
    sheet.append([])
    sheet.append(["", "", "", "Totals", credits, debits, ""])
    sheet.cell(row=sheet.max_row, column=4).font = Font(bold=True)
    for row in sheet.iter_rows(min_row=7):
        row[0].number_format = "DD/MM/YYYY"
        for cell in row[4:6]:
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
    for column, width in zip("ABCDEFG", (12, 18, 42, 34, 14, 14, 12)):
        sheet.column_dimensions[column].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _money(value):
    return f"{value:,.2f}" if value is not None else ""


def build_pdf(member, date_from=None, date_to=None):
    coop = CooperativeSettings.load()
    rows = statement_rows(member, date_from, date_to)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Statement {member.membership_number}", author=coop.name,
    )
    styles = getSampleStyleSheet()
    navy = colors.HexColor(f"#{NAVY}")
    styles["Title"].textColor = navy
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)

    story = [
        Paragraph(coop.name, styles["Title"]),
        Paragraph(coop.parent_institution, styles["BodyText"]),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>Member statement:</b> {member.full_name} ({member.membership_number})", styles["BodyText"]),
        Paragraph(f"<b>Period:</b> {_period_label(date_from, date_to)}", styles["BodyText"]),
        Paragraph(f"<b>Generated:</b> {timezone.localtime():%d %b %Y %H:%M}", styles["BodyText"]),
        Spacer(1, 5 * mm),
    ]
    # ReportLab's built-in fonts have no naira sign, so the PDF says NGN.
    data = [["Date", "Reference", "Description", "Account", "Credit (NGN)", "Debit (NGN)", "Status"]]
    for r in rows:
        data.append([
            f"{r['date']:%d/%m/%Y}", r["reference"], Paragraph(r["description"], small), Paragraph(r["account"], small),
            _money(r["credit"]), _money(r["debit"]), r["status"],
        ])
    credits, debits = _totals(rows)
    data.append(["", "", "", "Totals", _money(credits), _money(debits), ""])
    if not rows:
        data.insert(1, ["", "", "No transactions in this period.", "", "", "", ""])

    table = Table(data, colWidths=[22 * mm, 34 * mm, 80 * mm, 64 * mm, 26 * mm, 26 * mm, 20 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (4, 0), (5, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F5FB")]),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, navy),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    story.append(table)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("This statement is generated from the cooperative's ledger. Contact the secretariat about any discrepancy.", small))
    doc.build(story)
    return buffer.getvalue()
