"""Reading uploaded .xlsx / .csv tables (member imports, payroll batches)."""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

MAX_ROWS = 5000
DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"]


class SpreadsheetError(Exception):
    """The file as a whole can't be used (unreadable, empty, missing columns, too large)."""


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores long numbers as floats
    return str(value).strip()


def normalise_header(value):
    return re.sub(r"[^a-z0-9]+", "_", cell_text(value).lower()).strip("_")


def parse_date(value):
    """Cell -> date or None. Raises ValueError with a readable message."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = cell_text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError("Use a date like 2026-01-31 or 31/01/2026.")


def parse_amount(value):
    """Cell -> positive Decimal with at most 2 decimal places. Accepts '5,000.00' and '₦5000'."""
    raw = cell_text(value).replace(",", "").replace("₦", "").replace("NGN", "").strip()
    if not raw:
        raise ValueError("This field is required.")
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        raise ValueError("Enter an amount like 5000 or 5000.50.") from None
    if amount <= 0:
        raise ValueError("Amount must be greater than zero.")
    if amount != amount.quantize(Decimal("0.01")):
        raise ValueError("Amounts can have at most 2 decimal places (kobo).")
    return amount.quantize(Decimal("0.01"))


def read_table(uploaded_file, *, preferred_sheet=None, max_rows=MAX_ROWS, what="rows"):
    """Return (header_cells, data_rows) where data_rows is a list of (row_number, cells)."""
    name = (uploaded_file.name or "").lower()
    uploaded_file.seek(0)
    if name.endswith(".csv"):
        content = uploaded_file.read().decode("utf-8-sig", errors="replace")
        raw_rows = list(csv.reader(io.StringIO(content)))
    else:
        try:
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise SpreadsheetError("The file could not be read as an Excel workbook.") from exc
        if preferred_sheet and preferred_sheet in workbook.sheetnames:
            sheet = workbook[preferred_sheet]
        else:
            sheet = workbook.worksheets[0]
        raw_rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        workbook.close()

    numbered = [(i + 1, row) for i, row in enumerate(raw_rows) if any(cell_text(c) for c in row)]
    if not numbered:
        raise SpreadsheetError("The file is empty.")
    (_, header), data = numbered[0], numbered[1:]
    if not data:
        raise SpreadsheetError(f"The file has a header row but no {what}.")
    if len(data) > max_rows:
        raise SpreadsheetError(f"The file has {len(data)} rows; the maximum per upload is {max_rows}.")
    return header, data


def map_columns(header, aliases, required_labels):
    """
    Match header cells to column keys. aliases: {normalised header: key}.
    required_labels: {key: label} that must be present.
    Returns ({cell_index: key}, unknown_header_texts).
    """
    mapping, unknown = {}, []
    for index, cell in enumerate(header):
        key = aliases.get(normalise_header(cell))
        if key:
            mapping[index] = key
        elif cell_text(cell):
            unknown.append(cell_text(cell))
    missing = [label for key, label in required_labels.items() if key not in mapping.values()]
    if missing:
        raise SpreadsheetError(f"Missing required column(s): {', '.join(missing)}.")
    return mapping, unknown


_PERIOD_FORMATS = ["%Y-%m", "%m/%Y", "%b %Y", "%B %Y"]


def parse_period(value):
    """Cell -> first day of the month, or None. Accepts 2026-03, 03/2026, Mar 2026 or any date."""
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return (value.date() if isinstance(value, datetime) else value).replace(day=1)
    raw = cell_text(value)
    for fmt in _PERIOD_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date().replace(day=1)
        except ValueError:
            continue
    try:
        return parse_date(raw).replace(day=1)
    except ValueError:
        raise ValueError("Use a month like 2026-03 or Mar 2026.") from None


def build_template(sheet_title, columns, example_rows=(), notes=()):
    """
    Excel template: a data sheet with styled headers and example rows, plus an
    Instructions sheet. columns: [(label, required, help)].
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    fill = PatternFill("solid", fgColor="293C9C")
    for index, (label, required, _help) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=label + (" *" if required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        sheet.column_dimensions[cell.column_letter].width = max(14, len(label) + 4)
    for row in example_rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Column", "Required", "Notes"])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    for label, required, help_text in columns:
        instructions.append([label, "Yes" if required else "", help_text])
    if notes:
        instructions.append([])
        for note in notes:
            instructions.append([note])
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["C"].width = 70

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
