"""
Bulk member import from Excel (.xlsx) or CSV.

Flow: upload -> parse -> validate every row (dry run) -> officer reviews the
report -> commit. Commit is all-or-nothing and only allowed when no row has
errors, so the membership register never ends up half-imported.
"""
import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Callable

from django.contrib.auth import get_user_model
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db.models.functions import Upper
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.common.spreadsheets import SpreadsheetError, cell_text, map_columns, normalise_header, parse_date, read_table
from apps.configuration.models import Department

from .models import Member, MemberStatus

User = get_user_model()


# ---------------------------------------------------------------------------
# Cell parsers: return a clean value or raise ValueError with a readable message
# ---------------------------------------------------------------------------


def text(max_length):
    def parse(value):
        result = cell_text(value)
        if len(result) > max_length:
            raise ValueError(f"Must be at most {max_length} characters.")
        return result

    return parse


def identifier(max_length):
    def parse(value):
        return text(max_length)(value).upper()

    return parse


def phone(value):
    result = re.sub(r"[\s\-()]", "", cell_text(value))
    if not result:
        return ""
    if result.isdigit() and len(result) == 10:
        result = "0" + result  # leading zero dropped when Excel treats the number as numeric
    if not re.fullmatch(r"\+?\d{7,15}", result):
        raise ValueError("Enter a valid phone number, e.g. 08031234567 or +2348031234567.")
    return result


def email(value):
    result = cell_text(value).lower()
    if not result:
        return ""
    try:
        validate_email(result)
    except ValidationError:
        raise ValueError("Enter a valid email address.") from None
    return result


def past_date(value):
    result = parse_date(value)
    if result and result > date.today():
        raise ValueError("Cannot be in the future.")
    return result


def choice(choices, *, allowed=None):
    """Accept the stored value or its label, case-insensitively."""
    options = {}
    for stored, label in choices:
        if allowed is None or stored in allowed:
            options[str(stored).lower()] = stored
            options[str(label).lower()] = stored

    def parse(value):
        raw = cell_text(value).lower().rstrip(".")
        if not raw:
            return ""
        if raw not in options:
            names = sorted({label for stored, label in choices if allowed is None or stored in allowed})
            raise ValueError(f"Must be one of: {', '.join(names)}.")
        return options[raw]

    return parse


def nuban(value):
    result = cell_text(value)
    if not result:
        return ""
    if result.isdigit() and len(result) < 10:
        result = result.zfill(10)
    if not re.fullmatch(r"\d{10}", result):
        raise ValueError("Bank account numbers have exactly 10 digits.")
    return result


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    parse: Callable
    required: bool = False
    help: str = ""


COLUMNS = [
    Column("membership_number", "Membership number", identifier(30), help="Leave blank to generate one."),
    Column("title", "Title", choice(Member.Title.choices)),
    Column("first_name", "First name", text(100), required=True),
    Column("middle_name", "Middle name", text(100)),
    Column("last_name", "Last name", text(100), required=True),
    Column("gender", "Gender", choice(Member.Gender.choices)),
    Column("date_of_birth", "Date of birth", past_date, help="YYYY-MM-DD or DD/MM/YYYY"),
    Column("marital_status", "Marital status", choice(Member.MaritalStatus.choices)),
    Column("phone", "Phone", phone, required=True),
    Column("alt_phone", "Alternative phone", phone),
    Column("email", "Email", email, help="Needed for portal activation emails."),
    Column("residential_address", "Residential address", text(500)),
    Column("state_of_origin", "State of origin", text(50)),
    Column("lga", "LGA", text(100)),
    Column("staff_number", "Staff number", identifier(30)),
    Column("ippis_number", "IPPIS number", identifier(30)),
    Column("department", "Department", text(150), help="Department name or code, as set up in Settings."),
    Column("unit", "Unit", text(100)),
    Column("designation", "Designation", text(100)),
    Column("grade_level", "Grade level", text(20)),
    Column("employment_date", "Employment date", past_date),
    Column("employment_status", "Employment status", choice(Member.EmploymentStatus.choices)),
    Column("date_joined", "Date joined", past_date, help="Date the person joined the cooperative. Blank = today."),
    Column(
        "status",
        "Membership status",
        choice(MemberStatus.choices, allowed={MemberStatus.ACTIVE, MemberStatus.PENDING, MemberStatus.INACTIVE}),
        help="Active (default), Pending activation or Inactive.",
    ),
    Column("bank_name", "Bank name", text(100)),
    Column("bank_account_number", "Bank account number", nuban),
    Column("bank_account_name", "Bank account name", text(150)),
    Column("next_of_kin_name", "Next of kin name", text(200)),
    Column("next_of_kin_relationship", "Next of kin relationship", text(50)),
    Column("next_of_kin_phone", "Next of kin phone", phone),
]
COLUMNS_BY_KEY = {c.key: c for c in COLUMNS}
NEXT_OF_KIN_KEYS = ("next_of_kin_name", "next_of_kin_relationship", "next_of_kin_phone")


HEADER_ALIASES = {}
for _column in COLUMNS:
    HEADER_ALIASES[_column.key] = _column.key
    HEADER_ALIASES[normalise_header(_column.label)] = _column.key


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def _department_lookup():
    lookup = {}
    for dept in Department.objects.filter(is_active=True):
        lookup[dept.name.lower()] = dept
        if dept.code:
            lookup[dept.code.lower()] = dept
    return lookup


def _parse_rows(header, data):
    mapping, unknown = map_columns(header, HEADER_ALIASES, {c.key: c.label for c in COLUMNS if c.required})

    departments = _department_lookup()
    parsed = []
    for row_number, cells in data:
        values, errors = {}, {}
        for index, key in mapping.items():
            raw = cells[index] if index < len(cells) else None
            try:
                values[key] = COLUMNS_BY_KEY[key].parse(raw)
            except ValueError as exc:
                errors[key] = [str(exc)]
        for column in COLUMNS:
            values.setdefault(column.key, None if column.parse in (parse_date, past_date) else "")
            if column.required and not values[column.key] and column.key not in errors:
                errors[column.key] = ["This field is required."]

        if values["department"]:
            dept = departments.get(values["department"].lower())
            if dept is None:
                errors["department"] = [f"Unknown department '{values['department']}'."]
            else:
                values["department"] = str(dept.pk)

        kin = [values[k] for k in NEXT_OF_KIN_KEYS]
        if any(kin) and not all(kin):
            errors["next_of_kin_name"] = ["Give the next of kin's name, relationship and phone together."]

        values["status"] = values["status"] or MemberStatus.ACTIVE
        values["employment_status"] = values["employment_status"] or Member.EmploymentStatus.ACTIVE
        parsed.append({"row": row_number, "values": values, "errors": errors})
    return parsed, unknown


def _find_conflicts(parsed):
    """Duplicates within the file and against existing records. Mutates row error dicts."""
    unique_fields = {
        "email": "email address",
        "membership_number": "membership number",
        "staff_number": "staff number",
        "ippis_number": "IPPIS number",
    }
    seen = {field: {} for field in unique_fields}
    for row in parsed:
        for field, label in unique_fields.items():
            value = row["values"].get(field)
            if not value:
                continue
            if value in seen[field]:
                row["errors"].setdefault(field, []).append(
                    f"Duplicate {label}; also on row {seen[field][value]}."
                )
            else:
                seen[field][value] = row["row"]

    taken_emails = set(
        User.objects.filter(email__in=list(seen["email"]), member__isnull=False).values_list("email", flat=True)
    )
    taken = {"email": taken_emails}
    for field in ("membership_number", "staff_number", "ippis_number"):
        values = list(seen[field])
        taken[field] = set(
            Member.objects.annotate(_v=Upper(field)).filter(_v__in=values).values_list("_v", flat=True)
        ) if values else set()

    for row in parsed:
        for field, label in unique_fields.items():
            value = row["values"].get(field)
            if value and value in taken[field]:
                row["errors"].setdefault(field, []).append(f"An existing member already has this {label}.")


def _jsonable(values):
    return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in values.items()}


def validate_table(header, data):
    parsed, unknown = _parse_rows(header, data)
    _find_conflicts(parsed)
    errors = [{"row": r["row"], "errors": r["errors"]} for r in parsed if r["errors"]]
    report = {
        "total_rows": len(parsed),
        "valid_rows": len(parsed) - len(errors),
        "error_rows": len(errors),
        "errors": errors,
        "unknown_columns": unknown,
        "file_errors": [],
    }
    rows = [{"row": r["row"], "values": _jsonable(r["values"])} for r in parsed]
    return rows, report


def revalidate_stored_rows(rows):
    """Re-check stored rows against the database just before commit (records may have changed since upload)."""
    parsed = [{"row": r["row"], "values": dict(r["values"]), "errors": {}} for r in rows]
    _find_conflicts(parsed)
    return [{"row": r["row"], "errors": r["errors"]} for r in parsed if r["errors"]]


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

EXAMPLE_ROW = {
    "title": "Engr",
    "first_name": "Adaeze",
    "middle_name": "",
    "last_name": "Okafor",
    "gender": "Female",
    "date_of_birth": "1985-04-12",
    "phone": "08031234567",
    "email": "adaeze.okafor@example.com",
    "staff_number": "EMDI/0421",
    "ippis_number": "123456",
    "department": "",
    "designation": "Senior Engineer",
    "grade_level": "GL 12",
    "employment_date": "2012-03-01",
    "date_joined": "2013-01-15",
    "status": "Active",
    "next_of_kin_name": "Chidi Okafor",
    "next_of_kin_relationship": "Spouse",
    "next_of_kin_phone": "08039876543",
}


def build_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Members"
    header_fill = PatternFill("solid", fgColor="293C9C")
    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column.label + (" *" if column.required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        sheet.cell(row=2, column=index, value=EXAMPLE_ROW.get(column.key, ""))
        sheet.column_dimensions[cell.column_letter].width = max(14, len(column.label) + 4)
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Instructions")
    notes.append(["Column", "Required", "Notes"])
    for cell in notes[1]:
        cell.font = Font(bold=True)
    for column in COLUMNS:
        notes.append([column.label, "Yes" if column.required else "", column.help])
    notes.append([])
    notes.append(["Replace the example row with your members. Columns marked * are required."])
    notes.append(["Dates: YYYY-MM-DD or DD/MM/YYYY. Keep phone numbers as text so the leading 0 is kept."])
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["C"].width = 70

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_member_table(uploaded_file):
    return read_table(uploaded_file, preferred_sheet="Members", what="members")
