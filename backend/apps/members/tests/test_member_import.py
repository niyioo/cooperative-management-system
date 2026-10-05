import io

import pytest
from openpyxl import load_workbook

from apps.configuration.models import Department
from apps.members.models import Member, MemberImport
from apps.savings.models import SavingsAccount
from tests.factories import MemberFactory, make_officer

from .helpers import csv_file, xlsx_file

pytestmark = pytest.mark.django_db

IMPORTS = "/api/v1/admin/members/imports/"
HEADER = ["Membership number", "First name *", "Last name *", "Phone *", "Email", "Staff number", "Department",
          "Date joined", "Membership status", "Next of kin name", "Next of kin relationship", "Next of kin phone"]


def upload(client, file):
    return client.post(IMPORTS, {"file": file}, format="multipart")


def test_template_download(as_user, secretary):
    response = as_user(secretary).get(f"{IMPORTS}template/")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    headers = [c.value for c in workbook["Members"][1]]
    assert "First name *" in headers and "IPPIS number" in headers
    assert "Instructions" in workbook.sheetnames


def test_valid_import_creates_members(as_user, secretary):
    Department.objects.create(name="Metallurgy", code="MET")
    rows = [
        ["EMDI/OLD/0001", "Ada", "Obi", 8031234567, "ada@example.com", "emdi/1", "met", "15/01/2013", "Active",
         "Uche Obi", "Husband", "08030000001"],
        ["", "Musa", "Bello", "08029999999", "", "EMDI/2", "Metallurgy", "2019-06-01", "", "", "", ""],
    ]
    client = as_user(secretary)
    response = upload(client, xlsx_file(HEADER, rows))
    assert response.status_code == 201
    assert response.data["status"] == "READY", response.data["report"]
    assert response.data["total_rows"] == 2

    committed = client.post(f"{IMPORTS}{response.data['id']}/commit/", {"send_activation": False}, format="json")
    assert committed.status_code == 200
    assert committed.data["status"] == "COMMITTED"
    assert committed.data["created_count"] == 2

    ada = Member.objects.get(membership_number="EMDI/OLD/0001")
    assert ada.phone == "08031234567"  # leading zero restored
    assert ada.staff_number == "EMDI/1"
    assert ada.department.code == "MET"
    assert ada.status == Member.Status.ACTIVE
    assert ada.date_joined.isoformat() == "2013-01-15"
    assert ada.next_of_kin.get().full_name == "Uche Obi"
    musa = Member.objects.get(last_name="Bello")
    assert musa.membership_number.startswith("EMDI/COOP/")
    assert not musa.user.has_real_email
    assert SavingsAccount.objects.filter(member__in=[ada, musa], product__code="REGULAR").count() == 2
    assert MemberImport.objects.get().rows == []  # personal data not kept after commit


def test_row_errors_are_reported_and_block_commit(as_user, secretary):
    MemberFactory(user__email="taken@example.com")
    rows = [
        ["", "Ada", "", "08031234567", "", "S-1", "", "not a date", "", "", "", ""],
        ["", "Bola", "Ade", "08031234568", "taken@example.com", "S-1", "Nowhere", "", "", "Kemi", "", ""],
    ]
    client = as_user(secretary)
    response = upload(client, xlsx_file(HEADER, rows))
    assert response.data["status"] == "HAS_ERRORS"
    errors = {e["row"]: e["errors"] for e in response.data["report"]["errors"]}
    assert set(errors[2]) == {"last_name", "date_joined"}
    assert {"email", "staff_number", "department", "next_of_kin_name"} <= set(errors[3])
    assert "row 2" in errors[3]["staff_number"][0]

    blocked = client.post(f"{IMPORTS}{response.data['id']}/commit/", {}, format="json")
    assert blocked.data["error"]["code"] == "import_not_ready"
    assert not Member.objects.filter(first_name="Bola").exists()


def test_missing_required_column(as_user, secretary):
    response = upload(as_user(secretary), xlsx_file(["First name", "Phone"], [["Ada", "0803"]]))
    assert response.data["status"] == "HAS_ERRORS"
    assert "Last name" in response.data["report"]["file_errors"][0]


def test_csv_is_accepted(as_user, secretary):
    response = upload(as_user(secretary), csv_file(["first_name", "last_name", "phone"], [["Ada", "Obi", "08031234567"]]))
    assert response.data["status"] == "READY"


def test_conflicts_created_after_upload_are_caught_at_commit(as_user, secretary):
    client = as_user(secretary)
    response = upload(client, xlsx_file(["First name", "Last name", "Phone", "Staff number"], [["Ada", "Obi", "08031234567", "S-9"]]))
    MemberFactory(staff_number="S-9")

    committed = client.post(f"{IMPORTS}{response.data['id']}/commit/", {}, format="json")
    assert committed.data["error"]["code"] == "import_conflicts"
    assert MemberImport.objects.get(pk=response.data["id"]).status == MemberImport.Status.HAS_ERRORS


def test_import_needs_the_import_permission(as_user):
    treasurer = make_officer("Treasurer")
    assert as_user(treasurer).get(IMPORTS).status_code == 403
