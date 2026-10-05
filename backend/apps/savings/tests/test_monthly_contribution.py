"""The statutory monthly contribution, its arrears and the payroll deduction schedule (BR-29)."""
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.configuration.models import CooperativeSettings
from apps.ledger import corrections
from apps.members.models import Member
from apps.notifications.models import Notification
from apps.savings import services, statutory
from apps.savings.batches import ContributionBatchHandler
from apps.savings.models import MonthlyContributionChange, SavingsAccount
from apps.savings.statutory import add_months, this_month
from tests.factories import MemberFactory, make_officer

pytestmark = pytest.mark.django_db

MINE = "/api/v1/me/savings/monthly-contribution/"
ACCOUNTS = "/api/v1/admin/savings/accounts/"
SCHEDULE = "/api/v1/admin/savings/deduction-schedule/"
DOWNLOAD = "/api/v1/admin/savings/deduction-schedule/download/"


def months_ago(count):
    return add_months(this_month(), -count)


def ym(month):
    return month.strftime("%Y-%m")


@pytest.fixture
def regular(regular):
    regular.min_contribution = Decimal("5000.00")
    regular.save()
    return regular


@pytest.fixture
def tracked_from_six_months_ago(db):
    coop = CooperativeSettings.load()
    coop.contributions_tracked_from = months_ago(6)
    coop.save()
    return coop.contributions_tracked_from


@pytest.fixture
def saver(db, regular, tracked_from_six_months_ago):
    member = MemberFactory(date_joined=months_ago(24))
    SavingsAccount.objects.create(member=member, product=regular, opened_on=months_ago(24))
    return member


@pytest.fixture
def account(saver, regular):
    return SavingsAccount.objects.get(member=saver, product=regular)


def pay(officer, account, month, amount="5000"):
    return services.post_contribution(officer, account=account, amount=Decimal(amount), period=month, value_date=month)


class TestArrears:
    def test_nothing_paid_means_every_due_month_is_owed(self, account):
        p = statutory.position(account)
        assert p["amount"] == Decimal("5000.00")
        assert p["tracked_from"] == months_ago(6)
        assert p["months_due"] == 6  # the six months before this one; this month is not due yet
        assert p["arrears"] == Decimal("30000.00")
        assert p["months_behind"] == 6

    def test_paid_months_clear_and_a_missed_month_remains(self, account, accountant):
        for n in range(1, 7):
            if n != 3:
                pay(accountant, account, months_ago(n))
        p = statutory.position(account)
        assert p["arrears"] == Decimal("5000.00")
        assert p["months_behind"] == 1
        history = {h["month"]: h for h in p["history"]}
        assert history[ym(months_ago(3))]["paid"] == Decimal("0.00")
        assert history[ym(months_ago(1))]["paid"] == Decimal("5000.00")

    def test_paying_more_one_month_makes_up_for_another(self, account, accountant):
        for n in range(1, 7):
            if n != 3:
                pay(accountant, account, months_ago(n), "5000" if n != 1 else "10000")
        assert statutory.position(account)["arrears"] == Decimal("0.00")

    def test_this_month_counts_once_something_is_posted_for_it(self, account, accountant):
        pay(accountant, account, this_month(), "5000")
        p = statutory.position(account)
        assert p["months_due"] == 7
        assert p["paid_this_month"] == Decimal("5000.00")
        assert p["arrears"] == Decimal("30000.00")

    def test_a_reversed_contribution_no_longer_counts(self, account, accountant, treasurer):
        entry = pay(accountant, account, months_ago(1))
        assert statutory.position(account)["arrears"] == Decimal("25000.00")
        reversal = corrections.reverse_entry(accountant, entry, reason="Bounced")
        if reversal.status == "PENDING":
            from apps.ledger import services as ledger

            ledger.approve_entry(treasurer, reversal)
        assert statutory.position(account)["arrears"] == Decimal("30000.00")

    def test_tracking_starts_no_earlier_than_the_account(self, saver, account):
        account.opened_on = months_ago(2)
        account.save()
        assert statutory.position(account)["months_due"] == 2

    def test_without_a_tracking_month_the_account_opening_is_the_start(self, account):
        CooperativeSettings.objects.update(contributions_tracked_from=None)
        assert statutory.position(account)["months_due"] == 24

    def test_arrears_use_the_amount_that_applied_each_month(self, account):
        MonthlyContributionChange.objects.create(account=account, amount=Decimal("8000"), effective_from=months_ago(2))
        p = statutory.position(account)
        assert p["expected_total"] == Decimal("4") * 5000 + Decimal("2") * 8000
        assert p["amount"] == Decimal("8000.00")

    def test_frozen_accounts_and_inactive_members_build_no_arrears(self, saver, account):
        account.status = SavingsAccount.Status.FROZEN
        account.save()
        assert statutory.position(account)["arrears"] == 0
        account.status = SavingsAccount.Status.ACTIVE
        account.save()
        Member.objects.filter(pk=saver.pk).update(status=Member.Status.INACTIVE)
        assert statutory.position(account)["arrears"] == 0

    def test_arrears_report_lists_members_behind(self, as_user, treasurer, saver, account, accountant, regular):
        punctual = MemberFactory(date_joined=months_ago(24))
        on_time = SavingsAccount.objects.create(member=punctual, product=regular, opened_on=months_ago(24))
        for n in range(1, 7):
            pay(accountant, on_time, months_ago(n))
        response = as_user(treasurer).get("/api/v1/admin/reports/contribution-arrears/")
        assert response.status_code == 200, response.data
        rows = response.data["rows"]
        assert [r["membership_number"] for r in rows] == [saver.membership_number]
        assert rows[0]["arrears"] == "30000.00"

    def test_officer_dashboard_shows_the_arrears(self, as_user, treasurer, account):
        data = as_user(treasurer).get("/api/v1/admin/dashboard/").data
        assert data["savings"]["contribution_arrears"]["members"] == 1
        assert data["savings"]["contribution_arrears"]["amount"] == "30000.00"
        monthly = data["savings"]["monthly_contributions"]
        assert monthly["members"] == 1 and monthly["expected"] == "5000.00"
        assert monthly["collected"] == "0.00" and monthly["arrears_members"] == 1


class TestMemberChangesTheirAmount:
    def test_member_sees_their_position(self, as_user, saver, account):
        response = as_user(saver.user).get(MINE)
        assert response.status_code == 200
        assert response.data["amount"] == "5000.00"
        assert response.data["minimum"] == "5000.00"
        assert response.data["arrears"] == "30000.00"
        assert len(response.data["history"]) == 6

    def test_change_applies_from_next_month(self, as_user, saver, account):
        response = as_user(saver.user).post(MINE, {"amount": "12000"}, format="json")
        assert response.status_code == 200, response.data
        assert response.data["amount"] == "5000.00"  # this month is unchanged
        assert response.data["pending_change"] == {"amount": "12000.00", "effective_from": add_months(this_month(), 1)}
        account.refresh_from_db()
        assert account.elected_monthly_amount == Decimal("12000.00")
        log = AuditLog.objects.get(action="savings.monthly_contribution_changed")
        assert log.metadata["by_member"] is True

    def test_changing_again_replaces_the_pending_change(self, as_user, saver, account):
        client = as_user(saver.user)
        client.post(MINE, {"amount": "12000"}, format="json")
        client.post(MINE, {"amount": "9000"}, format="json")
        assert MonthlyContributionChange.objects.filter(account=account).count() == 1
        assert MonthlyContributionChange.objects.get(account=account).amount == Decimal("9000.00")

    def test_below_the_minimum_is_refused(self, as_user, saver, account):
        response = as_user(saver.user).post(MINE, {"amount": "4999.99"}, format="json")
        assert response.status_code == 400
        assert response.data["error"]["code"] == "below_minimum"
        assert "amount" in response.data["error"]["fields"]

    def test_unchanged_amount_is_refused(self, as_user, saver, account):
        response = as_user(saver.user).post(MINE, {"amount": "5000"}, format="json")
        assert response.data["error"]["code"] == "unchanged"

    def test_officers_have_no_member_endpoint(self, as_user, treasurer):
        assert as_user(treasurer).get(MINE).status_code == 403

    def test_dashboard_shows_the_monthly_contribution(self, as_user, saver, account):
        data = as_user(saver.user).get("/api/v1/me/dashboard/").data
        assert data["monthly_contribution"]["amount"] == "5000.00"
        assert data["monthly_contribution"]["arrears"] == "30000.00"


class TestOfficerChangesTheAmount:
    def url(self, account):
        return f"{ACCOUNTS}{account.pk}/monthly-contribution/"

    def test_officer_change_applies_this_month_and_tells_the_member(
        self, as_user, accountant, saver, account, mailoutbox, django_capture_on_commit_callbacks
    ):
        saver.user.email = "saver@example.org"
        saver.user.save()
        with django_capture_on_commit_callbacks(execute=True):
            response = as_user(accountant).post(
                self.url(account), {"amount": "7500", "reason": "Letter of 1 Oct"}, format="json"
            )
        assert response.status_code == 200, response.data
        assert response.data["amount"] == "7500.00"
        note = Notification.objects.get(recipient=saver.user)
        assert "7,500.00" in note.body and "Letter of 1 Oct" in note.body
        assert len(mailoutbox) == 1

    def test_officer_change_replaces_a_later_pending_change(self, as_user, accountant, saver, account):
        as_user(saver.user).post(MINE, {"amount": "12000"}, format="json")
        as_user(accountant).post(self.url(account), {"amount": "7500"}, format="json")
        p = statutory.position(account)
        assert p["amount"] == Decimal("7500.00")
        assert p["pending_change"] is None

    def test_officer_may_schedule_a_later_month_but_not_a_past_one(self, as_user, accountant, account):
        client = as_user(accountant)
        later = client.post(self.url(account), {"amount": "6000", "effective_from": ym(add_months(this_month(), 2))}, format="json")
        assert later.data["pending_change"]["amount"] == "6000.00"
        past = client.post(self.url(account), {"amount": "6000", "effective_from": ym(months_ago(1))}, format="json")
        assert past.data["error"]["code"] == "month_in_past"

    def test_viewing_needs_view_savings_and_changing_needs_posting_rights(self, as_user, secretary, account):
        assert as_user(secretary).get(self.url(account)).status_code == 403
        auditor = make_officer("Auditor")
        assert as_user(auditor).get(self.url(account)).status_code == 200
        assert as_user(auditor).post(self.url(account), {"amount": "6000"}, format="json").status_code == 403

    def test_christmas_accounts_have_no_monthly_contribution(self, as_user, accountant, saver, open_cycle):
        xmas = SavingsAccount.objects.get(member=saver, cycle=open_cycle)
        assert as_user(accountant).get(self.url(xmas)).status_code == 404

    def test_editing_the_account_routes_through_the_history(self, as_user, accountant, account):
        response = as_user(accountant).patch(f"{ACCOUNTS}{account.pk}/", {"elected_monthly_amount": "6500"}, format="json")
        assert response.status_code == 200, response.data
        assert MonthlyContributionChange.objects.get(account=account).effective_from == this_month()

    def test_opening_an_account_with_an_amount_records_it(self, accountant, regular, tracked_from_six_months_ago):
        member = MemberFactory()
        opened = services.open_account(accountant, member=member, product=regular, elected_monthly_amount=Decimal("8000"))
        assert statutory.position(opened)["amount"] == Decimal("8000.00")
        with pytest.raises(Exception):
            services.open_account(accountant, member=MemberFactory(), product=regular, elected_monthly_amount=Decimal("100"))


class TestMinimumChanges:
    def test_raising_the_minimum_keeps_past_months_and_lifts_low_amounts(self, super_admin, account, regular):
        services.update_product(super_admin, regular, min_contribution=Decimal("6000.00"))
        p = statutory.position(account)
        assert p["amount"] == Decimal("6000.00")
        assert p["expected_total"] == Decimal("30000.00")  # past months still at the old 5,000
        assert AuditLog.objects.filter(action="savings.contribution_minimum_changed").exists()

    def test_higher_amounts_are_left_alone(self, super_admin, account, regular):
        MonthlyContributionChange.objects.create(account=account, amount=Decimal("9000"), effective_from=months_ago(1))
        services.update_product(super_admin, regular, min_contribution=Decimal("6000.00"))
        assert statutory.position(account)["amount"] == Decimal("9000.00")


class TestDeductionSchedule:
    def test_schedule_adds_arrears_and_skips_members_already_recorded(self, as_user, treasurer, accountant, account, regular):
        paid_up = MemberFactory(date_joined=months_ago(24), staff_number="S-77", ippis_number="IP-77")
        done = SavingsAccount.objects.create(member=paid_up, product=regular, opened_on=months_ago(24))
        for n in range(0, 7):
            pay(accountant, done, months_ago(n))
        data = as_user(treasurer).get(SCHEDULE).data
        assert data["period"] == ym(this_month())
        assert data["totals"]["already_recorded"] == 1
        [row] = data["rows"]
        assert row["monthly"] == "5000.00" and row["arrears"] == "30000.00" and row["amount"] == "35000.00"

        without = as_user(treasurer).get(SCHEDULE, {"include_arrears": "false"}).data
        assert without["rows"][0]["amount"] == "5000.00"

    def test_a_future_month_uses_the_amount_that_will_apply(self, as_user, treasurer, saver, account):
        as_user(saver.user).post(MINE, {"amount": "12000"}, format="json")
        data = as_user(treasurer).get(SCHEDULE, {"period": ym(add_months(this_month(), 1))}).data
        assert data["rows"][0]["monthly"] == "12000.00"

    def test_past_months_never_add_arrears(self, as_user, treasurer, account):
        data = as_user(treasurer).get(SCHEDULE, {"period": ym(months_ago(2))}).data
        assert data["include_arrears"] is False
        assert data["rows"][0]["amount"] == "5000.00"

    def test_download_round_trips_through_the_contributions_batch(self, as_user, accountant, treasurer, saver, account):
        response = as_user(accountant).get(DOWNLOAD)
        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["Contributions", "Breakdown", "About"]
        sheet = workbook["Contributions"]
        assert [c.value for c in sheet[1]] == [
            "Membership number", "Staff number", "IPPIS number", "Name", "Product code", "Amount", "Month", "Reference",
        ]
        assert sheet.cell(row=2, column=1).value == saver.membership_number
        assert Decimal(str(sheet.cell(row=2, column=6).value)) == Decimal("35000")
        assert AuditLog.objects.filter(action="savings.deduction_schedule_exported").exists()

        upload = io.BytesIO(response.content)
        upload.name = "deductions.xlsx"
        from apps.common.spreadsheets import read_table

        header, rows = read_table(upload)
        lines, report = ContributionBatchHandler().validate_file(header, rows, {})
        assert report["errors"] == [] and report["unknown_columns"] == [], report
        assert len(lines) == 1

    def test_download_needs_posting_rights(self, as_user, account):
        auditor = make_officer("Auditor")
        assert as_user(auditor).get(SCHEDULE).status_code == 200
        assert as_user(auditor).get(DOWNLOAD).status_code == 403

    def test_no_statutory_product_is_a_clear_error(self, as_user, treasurer, regular):
        regular.is_mandatory = False
        regular.save()
        response = as_user(treasurer).get(SCHEDULE)
        assert response.data["error"]["code"] == "no_statutory_product"
