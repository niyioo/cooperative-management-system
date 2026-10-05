"""
The statutory monthly contribution (BR-29, ARCHITECTURE.md §6.3.1).

Every member has a monthly amount deducted from pay into the mandatory regular
savings product (Regular Savings for EMDI). The member chooses the amount; it
may not be below the product's minimum contribution. Members change it from the
next month; officers may change it from the current month.

Arrears are the amount expected over the months that are due, less what was
actually contributed for those months. A month is due once it has ended, or
as soon as anything is posted for it. Only posted contributions (net of
reversals) count, and only months from `contributions_tracked_from` (or the
account's opening) onwards.
"""
import math
from collections import defaultdict
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.db.models import Q, Sum
from django.utils import timezone

from apps.common.spreadsheets import workbook_bytes
from apps.accounts.permissions import assert_not_self, require_perm
from apps.accounts.perms import P
from apps.audit.services import record
from apps.common.exceptions import DomainError
from apps.configuration.models import CooperativeSettings
from apps.ledger.choices import TransactionStatus, TransactionType
from apps.ledger.models import Transaction
from apps.ledger.selectors import SIGNED_AMOUNT, ZERO
from apps.notifications import services as notifications
from apps.notifications.models import Notification

from .models import MonthlyContributionChange, ProductKind, SavingsAccount, SavingsProduct
from .rules import BLOCKED_MEMBER_STATUSES, month_start

HISTORY_MONTHS = 12
CONTRIBUTIONS = Q(txn_type=TransactionType.SAVINGS_CONTRIBUTION) | Q(
    txn_type=TransactionType.REVERSAL, reverses__txn_type=TransactionType.SAVINGS_CONTRIBUTION
)


def add_months(month, count):
    index = month.year * 12 + month.month - 1 + count
    return date(index // 12, index % 12 + 1, 1)


def this_month(today=None):
    return month_start(today or timezone.localdate())


def statutory_product():
    """The product monthly payroll deductions go into, or None if none is configured."""
    return (
        SavingsProduct.objects.filter(kind=ProductKind.REGULAR, is_mandatory=True, is_active=True)
        .order_by("display_order", "name")
        .first()
    )


def is_statutory(account, product=None):
    product = product or statutory_product()
    return product is not None and account.cycle_id is None and account.product_id == product.pk


def statutory_account(member, product=None):
    product = product or statutory_product()
    if product is None:
        return None
    return (
        SavingsAccount.objects.filter(member=member, product=product, cycle__isnull=True)
        .select_related("member", "product")
        .first()
    )


def _changes_by_account(account_ids):
    changes = defaultdict(list)
    for change in MonthlyContributionChange.objects.filter(account__in=account_ids).order_by("effective_from"):
        changes[change.account_id].append(change)
    return changes


def amount_for(changes, month, minimum):
    """The monthly amount that applies in `month`, given the account's changes in date order."""
    amount = None
    for change in changes:
        if change.effective_from > month:
            break
        amount = change.amount
    return minimum if amount is None else amount


def tracking_start(account, tracked_from):
    start = max(month_start(account.opened_on), month_start(account.member.date_joined))
    return max(start, tracked_from) if tracked_from else start


def _paid_by_month(account_ids, upto):
    rows = (
        Transaction.objects.posted()
        .filter(CONTRIBUTIONS, savings_account__in=account_ids, period__isnull=False, period__lte=upto)
        .values("savings_account", "period")
        .annotate(net=Sum(SIGNED_AMOUNT))
    )
    return {(r["savings_account"], r["period"]): r["net"] or ZERO for r in rows}


def positions(accounts, *, today=None, history=False):
    """
    {account_id: position} for statutory accounts: the amount this month, any
    change waiting to apply, and the arrears. Three queries whatever the count.
    """
    accounts = list(accounts)
    current = this_month(today)
    ids = [a.pk for a in accounts]
    changes = _changes_by_account(ids)
    paid = _paid_by_month(ids, current)
    tracked_from = CooperativeSettings.load().contributions_tracked_from
    result = {}
    for account in accounts:
        minimum = account.product.min_contribution
        own = changes.get(account.pk, [])
        tracked = account.status == SavingsAccount.Status.ACTIVE and account.member.status not in BLOCKED_MEMBER_STATUSES
        start = tracking_start(account, tracked_from)
        last_due = current if (account.pk, current) in paid else add_months(current, -1)

        months, month = [], start
        while month <= last_due:
            months.append({"month": month, "expected": amount_for(own, month, minimum), "paid": paid.get((account.pk, month), ZERO)})
            month = add_months(month, 1)
        expected_total = sum((m["expected"] for m in months), ZERO)
        paid_total = sum((m["paid"] for m in months), ZERO)
        arrears = max(expected_total - paid_total, ZERO) if tracked else ZERO

        amount = amount_for(own, current, minimum)
        pending = next((c for c in own if c.effective_from > current), None)
        entry = {
            "account_id": str(account.pk),
            "account_number": account.account_number,
            "product": account.product.name,
            "minimum": minimum,
            "maximum": account.product.max_monthly_contribution,
            "amount": amount,
            "pending_change": {"amount": pending.amount, "effective_from": pending.effective_from} if pending else None,
            "tracked": tracked,
            "tracked_from": start,
            "months_due": len(months),
            "expected_total": expected_total,
            "paid_total": paid_total,
            "arrears": arrears,
            "months_behind": math.ceil(arrears / amount) if arrears and amount else 0,
            "paid_this_month": paid.get((account.pk, current), ZERO),
        }
        if history:
            entry["history"] = [
                {"month": m["month"].strftime("%Y-%m"), "expected": m["expected"], "paid": m["paid"]}
                for m in reversed(months[-HISTORY_MONTHS:])
            ]
            entry["changes"] = [
                {"amount": c.amount, "effective_from": c.effective_from, "reason": c.reason, "created_at": c.created_at}
                for c in reversed(own)
            ]
        result[account.pk] = entry
    return result


def position(account, *, today=None):
    account = SavingsAccount.objects.select_related("member", "product").get(pk=account.pk)
    return positions([account], today=today, history=True)[account.pk]


def tracked_accounts(product=None):
    product = product or statutory_product()
    if product is None:
        return SavingsAccount.objects.none()
    return (
        SavingsAccount.objects.filter(product=product, cycle__isnull=True, status=SavingsAccount.Status.ACTIVE)
        .exclude(member__status__in=BLOCKED_MEMBER_STATUSES)
        .select_related("member", "member__department", "product")
        .order_by("member__membership_number")
    )


def month_summary(today=None):
    """This month's monthly-contribution figures for the officer dashboard, plus who is behind."""
    accounts = list(tracked_accounts())
    found = positions(accounts, today=today)
    behind = sorted(((a, found[a.pk]) for a in accounts if found[a.pk]["arrears"] > 0), key=lambda pair: pair[1]["arrears"], reverse=True)
    return {
        "month": this_month(today).strftime("%Y-%m"),
        "members": len(accounts),
        "expected": sum((p["amount"] for p in found.values()), ZERO),
        "collected": sum((p["paid_this_month"] for p in found.values()), ZERO),
        "members_paid": sum(1 for p in found.values() if p["paid_this_month"] > 0),
        "arrears_members": len(behind),
        "arrears_amount": sum((p["arrears"] for _, p in behind), ZERO),
        "behind": behind,
    }


def arrears_rows(department=None):
    """Every member behind on their monthly contribution, largest arrears first."""
    accounts = tracked_accounts()
    if department:
        accounts = accounts.filter(member__department_id=department)
    accounts = list(accounts)
    found = positions(accounts)
    rows = [(a, found[a.pk]) for a in accounts if found[a.pk]["arrears"] > 0]
    rows.sort(key=lambda pair: pair[1]["arrears"], reverse=True)
    return rows


# ---------------------------------------------------------------------------
# Changing the amount
# ---------------------------------------------------------------------------

def _field_error(field, message, code):
    raise DomainError(message, code=code, fields={field: [message]})


def amount_problem(product, amount):
    """(message, code) if `amount` is outside the product's monthly limits, else None."""
    if amount < product.min_contribution:
        return f"The minimum monthly contribution is ₦{product.min_contribution:,.2f}.", "below_minimum"
    if product.max_monthly_contribution is not None and amount > product.max_monthly_contribution:
        return f"The maximum monthly contribution is ₦{product.max_monthly_contribution:,.2f}.", "above_maximum"
    return None


@transaction.atomic
def set_monthly_contribution(actor, account, *, amount, effective_from=None, reason="", by_member=False):
    """
    Change the monthly amount. Members (by_member=True, their own account only)
    change it from next month, since this month's deduction may already be with
    payroll. Officers may choose the current month or any later one. A later
    change already scheduled is replaced.
    """
    product = statutory_product()
    if product is None or not is_statutory(account, product):
        raise DomainError("This is not the monthly contribution account.", code="not_statutory_account")
    account = SavingsAccount.objects.select_for_update(of=("self",)).select_related("member", "product").get(pk=account.pk)
    member = account.member
    if by_member:
        if member.user_id != actor.pk:
            raise DomainError("You can only change your own monthly contribution.", code="not_your_account")
    else:
        require_perm(actor, P.POST_SAVINGS_CONTRIBUTION)
        assert_not_self(actor, member, "change the monthly contribution of")
    if account.status != SavingsAccount.Status.ACTIVE:
        raise DomainError(f"The {product.name} account is {account.get_status_display().lower()}.", code="account_not_active")
    if member.status in BLOCKED_MEMBER_STATUSES:
        raise DomainError(f"{member.get_status_display()} members have no monthly contribution.", code="member_not_active")

    amount = Decimal(amount).quantize(Decimal("0.01"))
    if amount <= 0:
        _field_error("amount", "Enter an amount greater than zero.", "invalid_amount")
    problem = amount_problem(product, amount)
    if problem:
        _field_error("amount", *problem)

    current = this_month()
    if by_member:
        effective = add_months(current, 1)
    else:
        effective = month_start(effective_from) if effective_from else current
        if effective < current:
            _field_error("effective_from", "A change can only apply from this month onwards.", "month_in_past")

    changes = list(account.contribution_changes.order_by("effective_from"))
    before = amount_for(changes, effective, product.min_contribution)
    superseded = [c for c in changes if c.effective_from > effective]
    if amount == before and not superseded:
        _field_error("amount", f"The monthly contribution from {effective:%B %Y} is already ₦{amount:,.2f}.", "unchanged")

    for change in superseded:
        change.delete()  # a plan that never took effect; the audit entry keeps it
    change, _ = MonthlyContributionChange.objects.update_or_create(
        account=account,
        effective_from=effective,
        defaults={"amount": amount, "changed_by": actor, "reason": (reason or "").strip()[:255]},
    )
    if account.elected_monthly_amount != amount:
        account.elected_monthly_amount = amount
        account.save(update_fields=["elected_monthly_amount", "updated_at"])

    record(
        "savings.monthly_contribution_changed",
        actor=actor,
        obj=account,
        changes={"monthly_contribution": [before, amount]},
        metadata={
            "member": member.membership_number,
            "effective_from": effective,
            "reason": change.reason,
            "by_member": by_member,
            "superseded": [{"amount": c.amount, "effective_from": c.effective_from} for c in superseded],
        },
    )
    if not by_member:
        notifications.notify_member(
            member,
            category=Notification.Category.SAVINGS,
            title="Your monthly contribution has changed",
            body=(
                f"Your monthly {product.name} contribution is now {notifications.naira(amount)} "
                f"from {effective:%B %Y}." + (f" Reason: {change.reason}" if change.reason else "")
            ),
            link=notifications.link("savings"),
            email=True,
        )
    return change


def minimum_changed(actor, product, old_minimum):
    """
    The product minimum moved. Months already past keep the amounts that applied
    then: accounts with no recorded amount get the old minimum written down
    from their start. From this month, anyone below the new minimum moves up to it.
    """
    if not product.is_mandatory or product.kind != ProductKind.REGULAR:
        return 0
    current = this_month()
    accounts = list(SavingsAccount.objects.filter(product=product, cycle__isnull=True).exclude(status=SavingsAccount.Status.CLOSED).select_related("member"))
    changes = _changes_by_account([a.pk for a in accounts])
    tracked_from = CooperativeSettings.load().contributions_tracked_from
    raised = 0
    for account in accounts:
        own = changes.get(account.pk, [])
        start = tracking_start(account, tracked_from)
        if not own and start < current:
            MonthlyContributionChange.objects.create(
                account=account, amount=old_minimum, effective_from=start, changed_by=actor, reason="Minimum at the time"
            )
        if product.min_contribution > old_minimum:
            for change in own:
                if change.effective_from > current and change.amount < product.min_contribution:
                    change.amount = product.min_contribution
                    change.save(update_fields=["amount", "updated_at"])
            if amount_for(own, current, old_minimum) < product.min_contribution:
                MonthlyContributionChange.objects.update_or_create(
                    account=account,
                    effective_from=current,
                    defaults={"amount": product.min_contribution, "changed_by": actor, "reason": "Raised to the new minimum"},
                )
                account.elected_monthly_amount = product.min_contribution
                account.save(update_fields=["elected_monthly_amount", "updated_at"])
                raised += 1
    record(
        "savings.contribution_minimum_changed",
        actor=actor,
        obj=product,
        changes={"min_contribution": [old_minimum, product.min_contribution]},
        metadata={"accounts_raised": raised},
    )
    return raised


# ---------------------------------------------------------------------------
# Payroll deduction schedule
# ---------------------------------------------------------------------------

def deduction_schedule(period, *, include_arrears=True, today=None):
    """
    What payroll (IPPIS) should deduct for `period`: each member's monthly
    amount, plus arrears when asked and the period is not already past.
    Members whose contribution for the period is already recorded are left out.
    The rows match the CONTRIBUTIONS batch upload, so the same file comes back.
    """
    product = statutory_product()
    if product is None:
        raise DomainError("No mandatory regular savings product is set up for monthly deductions.", code="no_statutory_product")
    period = month_start(period)
    current = this_month(today)
    add_arrears = include_arrears and period >= current
    accounts = [a for a in tracked_accounts(product) if month_start(a.opened_on) <= period]
    ids = [a.pk for a in accounts]
    recorded = set(
        Transaction.objects.filter(
            savings_account__in=ids,
            period=period,
            txn_type=TransactionType.SAVINGS_CONTRIBUTION,
            status__in=[TransactionStatus.POSTED, TransactionStatus.PENDING],
        ).values_list("savings_account", flat=True)
    )
    changes = _changes_by_account(ids)
    found = positions(accounts, today=today) if add_arrears else {}

    rows, already = [], 0
    for account in accounts:
        if account.pk in recorded:
            already += 1
            continue
        monthly = amount_for(changes.get(account.pk, []), period, product.min_contribution)
        arrears = found[account.pk]["arrears"] if add_arrears else ZERO
        if monthly + arrears <= 0:
            continue
        member = account.member
        rows.append({
            "account_id": str(account.pk),
            "member_id": str(member.pk),
            "membership_number": member.membership_number,
            "staff_number": member.staff_number or "",
            "ippis_number": member.ippis_number or "",
            "name": member.full_name,
            "department": member.department.name if member.department_id else "",
            "monthly": monthly,
            "arrears": arrears,
            "amount": monthly + arrears,
        })
    return {
        "period": period.strftime("%Y-%m"),
        "product": {"id": str(product.pk), "code": product.code, "name": product.name, "minimum": product.min_contribution},
        "include_arrears": add_arrears,
        "reference": f"PAYROLL-{period:%Y-%m}",
        "rows": rows,
        "totals": {
            "members": len(rows),
            "monthly": sum((r["monthly"] for r in rows), ZERO),
            "arrears": sum((r["arrears"] for r in rows), ZERO),
            "amount": sum((r["amount"] for r in rows), ZERO),
            "already_recorded": already,
        },
    }


def schedule_workbook(schedule, cooperative_name):
    """The schedule as .xlsx: the upload sheet first, then a breakdown and notes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="293C9C")

    def headed(sheet, labels, widths):
        sheet.append(labels)
        for cell, width in zip(sheet[1], widths):
            cell.font, cell.fill = header_font, header_fill
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.freeze_panes = "A2"

    workbook = Workbook()
    upload = workbook.active
    upload.title = "Contributions"
    headed(
        upload,
        ["Membership number", "Staff number", "IPPIS number", "Name", "Product code", "Amount", "Month", "Reference"],
        [20, 14, 14, 30, 14, 14, 10, 20],
    )
    code, period, reference = schedule["product"]["code"], schedule["period"], schedule["reference"]
    for r in schedule["rows"]:
        upload.append([r["membership_number"], r["staff_number"], r["ippis_number"], r["name"], code, r["amount"], period, reference])
        upload.cell(row=upload.max_row, column=6).number_format = "#,##0.00"

    breakdown = workbook.create_sheet("Breakdown")
    headed(breakdown, ["Membership number", "Name", "Department", "Monthly contribution", "Arrears", "Total deduction"], [20, 30, 24, 20, 14, 16])
    for r in schedule["rows"]:
        breakdown.append([r["membership_number"], r["name"], r["department"], r["monthly"], r["arrears"], r["amount"]])
    totals = schedule["totals"]
    breakdown.append(["Total", f"{totals['members']} members", "", totals["monthly"], totals["arrears"], totals["amount"]])
    for row in breakdown.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"
    for cell in breakdown[breakdown.max_row]:
        cell.font = Font(bold=True)

    about = workbook.create_sheet("About")
    month_label = date.fromisoformat(f"{period}-01").strftime("%B %Y")
    for line in [
        [f"{cooperative_name}: monthly contribution deduction schedule"],
        ["Month", month_label],
        ["Savings product", f"{schedule['product']['name']} ({code})"],
        ["Members", totals["members"]],
        ["Total to deduct", totals["amount"]],
        ["Arrears included", "Yes" if schedule["include_arrears"] else "No"],
        ["Already recorded for this month (left out)", totals["already_recorded"]],
        ["Generated", timezone.localtime().strftime("%d %b %Y %H:%M")],
        [],
        ["Send the Contributions sheet to payroll. Once the deductions are made, upload the same sheet under"],
        ["Batches as a Contributions batch (correct any amount that payroll could not deduct first)."],
    ]:
        about.append(line)
    about.column_dimensions["A"].width = 44
    about.column_dimensions["B"].width = 30
    about["A1"].font = Font(bold=True, size=13)

    return workbook_bytes(workbook)
