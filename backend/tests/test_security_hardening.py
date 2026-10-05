"""Rate limits on the whole API and on password changes, no caching of API responses, and the monthly contribution ceiling."""
from decimal import Decimal

import pytest
from django.core.cache import cache

from apps.accounts.throttles import APIAnonRateThrottle, APIUserRateThrottle, PasswordChangeRateThrottle
from apps.savings import services
from apps.savings.models import SavingsAccount, SavingsProduct

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def fresh_counters():
    cache.clear()
    yield
    cache.clear()


class TestRateLimits:
    def test_signed_in_requests_are_limited_per_user(self, as_user, member, monkeypatch):
        monkeypatch.setattr(APIUserRateThrottle, "THROTTLE_RATES", {"user": "2/min"})
        client = as_user(member.user)
        assert client.get("/api/v1/me/dashboard/").status_code == 200
        assert client.get("/api/v1/me/savings/").status_code == 200
        response = client.get("/api/v1/me/dashboard/")
        assert response.status_code == 429
        assert response.data["error"]["code"] == "throttled"

    def test_signed_out_requests_are_limited_per_address(self, api, monkeypatch):
        monkeypatch.setattr(APIAnonRateThrottle, "THROTTLE_RATES", {"anon": "2/min"})
        refresh = "/api/v1/auth/token/refresh/"
        api.post(refresh, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        api.post(refresh, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        assert api.post(refresh, HTTP_X_REQUESTED_WITH="XMLHttpRequest").status_code == 429

    def test_health_checks_are_never_limited(self, api, monkeypatch):
        monkeypatch.setattr(APIAnonRateThrottle, "THROTTLE_RATES", {"anon": "1/min"})
        assert all(api.get("/api/v1/health/").status_code == 200 for _ in range(3))

    def test_guessing_the_current_password_is_limited(self, as_user, member, monkeypatch):
        monkeypatch.setattr(PasswordChangeRateThrottle, "THROTTLE_RATES", {"password_change": "2/min"})
        client = as_user(member.user)
        attempt = {"current_password": "wrong-guess-1", "new_password": "A-new-password-2026"}
        client.post("/api/v1/auth/password/change/", attempt, format="json")
        client.post("/api/v1/auth/password/change/", attempt, format="json")
        assert client.post("/api/v1/auth/password/change/", attempt, format="json").status_code == 429


class TestNoCaching:
    def test_login_response_with_the_access_token_is_never_stored(self, api, member, password):
        response = api.post("/api/v1/auth/login/", {"identifier": member.user.email, "password": password}, format="json")
        assert response.status_code == 200 and "access" in response.data
        assert response["Cache-Control"] == "no-store"
        assert response["Pragma"] == "no-cache"

    def test_financial_data_is_never_stored(self, as_user, member):
        assert as_user(member.user).get("/api/v1/me/savings/")["Cache-Control"] == "no-store"

    def test_refresh_token_is_only_in_an_httponly_cookie(self, api, member, password):
        response = api.post("/api/v1/auth/login/", {"identifier": member.user.email, "password": password}, format="json")
        assert "refresh" not in response.data
        cookie = next(iter(response.cookies.values()))
        assert cookie["httponly"] and cookie["path"] == "/api/v1/auth/"


class TestMonthlyContributionCeiling:
    @pytest.fixture
    def regular(self, db):
        return SavingsProduct.objects.get(code="REGULAR")

    def test_emdi_defaults_are_a_5000_minimum_and_no_maximum(self, regular):
        assert regular.min_contribution == Decimal("5000.00")
        assert regular.expected_monthly_contribution == Decimal("5000.00")
        assert regular.max_monthly_contribution is None

    def test_member_cannot_choose_more_than_the_maximum(self, as_user, member, super_admin, regular):
        services.update_product(super_admin, regular, max_monthly_contribution=Decimal("1000000"))
        SavingsAccount.objects.get_or_create(member=member, product=regular, cycle=None)
        client = as_user(member.user)
        response = client.post("/api/v1/me/savings/monthly-contribution/", {"amount": "1000000.01"}, format="json")
        assert response.status_code == 400
        assert response.data["error"]["code"] == "above_maximum"
        assert client.get("/api/v1/me/savings/monthly-contribution/").data["maximum"] == "1000000.00"
        assert client.post("/api/v1/me/savings/monthly-contribution/", {"amount": "1000000"}, format="json").status_code == 200

    def test_maximum_cannot_be_below_the_minimum(self, super_admin, regular):
        from apps.common.exceptions import DomainError

        with pytest.raises(DomainError) as raised:
            services.update_product(super_admin, regular, max_monthly_contribution=Decimal("4000"))
        assert raised.value.code == "invalid_maximum"

    def test_blank_maximum_means_no_limit(self, as_user, member, regular):
        SavingsAccount.objects.get_or_create(member=member, product=regular, cycle=None)
        response = as_user(member.user).post("/api/v1/me/savings/monthly-contribution/", {"amount": "5000000"}, format="json")
        assert response.status_code == 200


class TestSpreadsheetFormulaInjection:
    def test_text_that_looks_like_a_formula_is_exported_as_plain_text(self):
        import io

        from openpyxl import Workbook, load_workbook

        from apps.common.spreadsheets import workbook_bytes

        workbook = Workbook()
        workbook.active.append(['=HYPERLINK("http://example.invalid","Click")', "+1+2", "@SUM(A1)", "Ada", 5000])
        cells = load_workbook(io.BytesIO(workbook_bytes(workbook))).active[1]
        assert [c.data_type for c in cells] == ["s", "s", "s", "s", "n"]
        assert cells[0].value.startswith("=HYPERLINK")  # kept, but as text

    def test_member_statement_export_keeps_descriptions_as_text(self, as_user, member):
        import io

        from django.utils import timezone
        from openpyxl import load_workbook

        from apps.ledger.models import Transaction
        from apps.savings.models import SavingsAccount, SavingsProduct

        account, _ = SavingsAccount.objects.get_or_create(member=member, product=SavingsProduct.objects.get(code="REGULAR"), cycle=None)
        Transaction.objects.create(member=member, txn_type="SAVINGS_OPENING_BALANCE", entry_side="CREDIT", amount=Decimal("100"),
                                   savings_account=account, status="POSTED", posted_at=timezone.now(), description="=1+1")
        response = as_user(member.user).get("/api/v1/me/transactions/statement/", {"format": "xlsx"})
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert all(c.data_type != "f" for row in sheet.iter_rows() for c in row)


class TestSupportConsoleSignIn:
    def test_repeated_failures_lock_the_address_out_and_are_audited(self, client, super_admin):
        from django.conf import settings

        from apps.audit.models import AuditLog

        url = f"/{settings.DJANGO_ADMIN_URL}login/"
        for _ in range(5):
            assert client.post(url, {"username": super_admin.email, "password": "wrong"}).status_code == 200
        assert client.post(url, {"username": super_admin.email, "password": "wrong"}).status_code == 429
        assert AuditLog.objects.filter(action="auth.admin_login_failed").count() == 5


class TestPasswordChangeEndsSessions:
    def test_access_tokens_issued_before_a_password_change_stop_working(self, as_user, member):
        import datetime

        from django.utils import timezone

        client = as_user(member.user)
        assert client.get("/api/v1/me/dashboard/").status_code == 200
        type(member.user).objects.filter(pk=member.user.pk).update(last_password_change=timezone.now() + datetime.timedelta(seconds=5))
        response = client.get("/api/v1/me/dashboard/")
        assert response.status_code == 401

    def test_a_fresh_sign_in_after_the_change_works(self, api, member, password):
        response = api.post("/api/v1/auth/login/", {"identifier": member.user.email, "password": password}, format="json")
        api.credentials(HTTP_AUTHORIZATION=f"Bearer {response.data['access']}")
        assert api.get("/api/v1/me/dashboard/").status_code == 200
